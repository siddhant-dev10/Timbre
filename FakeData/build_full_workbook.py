#!/usr/bin/env python3
"""
Timbre Project - Master Synthetic Demo Dataset Generator
Generates:
  - Timbre_Fake_Citizens.xlsx (Professional Excel Workbook with 4 Sheets)
  - Timbre_Fake_Citizens.csv (CSV export of Sheet 1: Citizens)
"""

import os
import csv
import shutil
from datetime import date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from make_records_1_25 import RECORDS_1_25
from make_records_26_50 import RECORDS_26_50
from make_records_51_75 import RECORDS_51_75
from make_records_76_100 import RECORDS_76_100

ALL_RECORDS = RECORDS_1_25 + RECORDS_26_50 + RECORDS_51_75 + RECORDS_76_100
assert len(ALL_RECORDS) == 100, f"Expected 100 records, got {len(ALL_RECORDS)}"

REF_DATE = date(2026, 10, 27)

def calc_age(dob_str: str, ref: date = REF_DATE) -> int:
    parts = [int(p) for p in dob_str.split('-')]
    dob = date(parts[0], parts[1], parts[2])
    age = ref.year - dob.year
    if (ref.month, ref.day) < (dob.month, dob.day):
        age -= 1
    return age

# -------------------------------------------------------------
# 1. CATALOGUE OF VERIFIED OFFICIAL GOVERNMENT SERVICES (18 SCHEMES)
# -------------------------------------------------------------
SERVICES_DATA = [
    {
        "service_id": "VOTER_ID_ECI",
        "service_name": "Voter Registration (Form 6)",
        "service_category": "Civic Identity & Electoral",
        "intended_beneficiaries": "Indian citizens who have attained or are approaching 18 years of age",
        "basic_eligibility": "Must be an Indian citizen, aged 18 or above on the qualifying date (or advance application at 17+), ordinary resident of the polling area",
        "commonly_required_documents": "Recent passport photograph, Proof of Age (Birth Certificate / Aadhaar / 10th marksheet), Proof of Ordinary Residence (Electricity Bill / Bank Passbook / Ration Card)",
        "official_application_portal": "Voters' Service Portal",
        "official_source_url": "https://voters.eci.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Statutory electoral registration administered by Election Commission of India. ₹0 statutory fee. ECI accepts advance Form 6 applications for citizens turning 18 in upcoming quarterly cycles."
    },
    {
        "service_id": "LEARNERS_LICENCE_MORTH",
        "service_name": "Learner's Driving Licence (Sarathi)",
        "service_category": "Transport & Licensing",
        "intended_beneficiaries": "Individuals seeking preliminary driving authorization (age 16+ for gearless <50cc with consent, age 18+ for Light Motor Vehicles)",
        "basic_eligibility": "Age >= 18 for LMV (cars/motorcycles with gear); physical fitness declaration Form 1/1A; resident of the RTO jurisdiction",
        "commonly_required_documents": "Aadhaar (for contactless e-KYC test), Age proof (Birth Certificate / 10th certificate), Address proof, Form 1 Self-declaration",
        "official_application_portal": "Parivahan Sewa / Sarathi Portal",
        "official_source_url": "https://parivahan.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by MoRTH and State Transport Departments. Statutory online application fee ~₹150-₹200. Contactless Aadhaar-based test allows home-based testing in participating states."
    },
    {
        "service_id": "AYUSHMAN_BHARAT_PMJAY",
        "service_name": "Pradhan Mantri Jan Arogya Yojana (AB PM-JAY)",
        "service_category": "Healthcare & Social Protection",
        "intended_beneficiaries": "Bottom 40% economically vulnerable families identified through SECC 2011 deprivation criteria and NFSA ration card databases",
        "basic_eligibility": "Family listed in SECC 2011 deprivation database or priority NFSA ration card; no formal/government employee in family; no income tax payees",
        "commonly_required_documents": "Ration Card (NFSA / BPL / Antyodaya), Aadhaar Card of all family members, Active mobile number",
        "official_application_portal": "PM-JAY Beneficiary Portal",
        "official_source_url": "https://pmjay.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by National Health Authority (NHA). Provides ₹5,00,000 cashless secondary and tertiary hospitalisation cover per family per year at empaneled hospitals. Card issuance is ₹0."
    },
    {
        "service_id": "SENIOR_CITIZEN_PMJAY_70",
        "service_name": "Ayushman Vaya Vandana Card (PM-JAY 70+ Universal Cover)",
        "service_category": "Healthcare & Geriatric Welfare",
        "intended_beneficiaries": "All Indian senior citizens aged 70 years and above, irrespective of family income or socioeconomic status",
        "basic_eligibility": "Attainment of age 70 or above (verified by Aadhaar date of birth); Indian citizen; individual enrollment",
        "commonly_required_documents": "Aadhaar Card with active mobile link (for OTP) or biometric verification at CSC/hospital",
        "official_application_portal": "NHA Beneficiary Portal & Ayushman App",
        "official_source_url": "https://beneficiary.nha.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Universal health coverage of ₹5 lakh per year for all citizens 70+. Provides dedicated card distinct from regular family PM-JAY. Free enrollment at beneficiary.nha.gov.in."
    },
    {
        "service_id": "PM_KISAN",
        "service_name": "Pradhan Mantri Kisan Samman Nidhi (PM-KISAN)",
        "service_category": "Agriculture & Farmer Support",
        "intended_beneficiaries": "Landholding farmer families owning cultivable agricultural land across India",
        "basic_eligibility": "Farmer family owning cultivable land in state land records; exclusions apply: institutional landholders, income tax payees, constitutional post holders, retired government staff",
        "commonly_required_documents": "Aadhaar Card, Land revenue record (Khatauni / Jamabandi / RoR extract), Active bank account passbook (Aadhaar-NPCI seeded)",
        "official_application_portal": "PM-KISAN Portal",
        "official_source_url": "https://pmkisan.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Ministry of Agriculture & Farmers Welfare. ₹6,000 per year delivered in 3 four-monthly installments of ₹2,000 directly via DBT. Mandatory e-KYC via facial recognition or OTP."
    },
    {
        "service_id": "KCC",
        "service_name": "Kisan Credit Card (KCC) Scheme",
        "service_category": "Agriculture & Rural Finance",
        "intended_beneficiaries": "Owner cultivators, tenant farmers, oral lessees, sharecroppers, animal husbandry and fisheries farmers",
        "basic_eligibility": "Engaged in crop cultivation or animal husbandry/fisheries; age between 18 and 75; verifiable land revenue record or tenancy agreement",
        "commonly_required_documents": "One-page KCC application form, Aadhaar Card / Voter ID, Land ownership record / village patwari crop certificate, Passport photographs",
        "official_application_portal": "PM-KISAN / National Portal of India",
        "official_source_url": "https://pmkisan.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by NABARD, RBI, and commercial/cooperative banks. Short-term crop loans up to ₹3 lakh at subsidized 7% interest (effective 4% with 3% prompt repayment incentive)."
    },
    {
        "service_id": "PM_AWAS_YOJANA_G",
        "service_name": "Pradhan Mantri Awaas Yojana - Gramin (PMAY-G)",
        "service_category": "Housing & Rural Development",
        "intended_beneficiaries": "Homeless rural households and those living in kutcha or severely dilapidated houses in rural areas",
        "basic_eligibility": "Listed in SECC 2011 deprivation list or verified Awaas+ survey; no pucca house owned by any family member anywhere in India",
        "commonly_required_documents": "Aadhaar Card, MGNREGA Job Card number, Bank account passbook, Land possession certificate / Gram Panchayat NOC",
        "official_application_portal": "AwaasSoft PMAY-G Portal",
        "official_source_url": "https://pmayg.nic.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Ministry of Rural Development. Direct financial grant of ₹1,20,000 (plains) / ₹1,30,000 (hilly states) plus 90-95 days of unskilled labour under MGNREGA and ₹12,000 for toilet construction under SBM."
    },
    {
        "service_id": "PMMVY",
        "service_name": "Pradhan Mantri Matru Vandana Yojana (PMMVY)",
        "service_category": "Maternal & Child Health",
        "intended_beneficiaries": "Pregnant women and lactating mothers for the first live birth (and second child if girl)",
        "basic_eligibility": "Age 19 years or older; belonging to socially/economically disadvantaged categories (SC/ST, PwD, BPL/EWS ration card, PM-JAY beneficiary, e-Shram registered); not in regular government employment",
        "commonly_required_documents": "Mother and Father Aadhaar Cards, Mother and Child Protection (MCP) Card, Bank passbook linked to Aadhaar, Proof of qualifying category",
        "official_application_portal": "PMMVY Portal (Ministry of WCD)",
        "official_source_url": "https://pmmvy.wcd.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Ministry of Women and Child Development. Cash DBT incentive: ₹5,000 in two installments for first child; ₹6,000 in single installment for second girl child."
    },
    {
        "service_id": "SUKANYA_SAMRIDDHI",
        "service_name": "Sukanya Samriddhi Account (Beti Bachao Beti Padhao)",
        "service_category": "Girl Child Savings & Welfare",
        "intended_beneficiaries": "Parents and legal guardians of girl children aged 10 years or younger",
        "basic_eligibility": "Girl child must be an Indian resident, age 10 or below at account opening; max 2 accounts per family (except biological twins/triplets)",
        "commonly_required_documents": "Birth Certificate of the girl child, Identity and address proof of parent/guardian (Aadhaar/PAN), Recent passport photographs",
        "official_application_portal": "India Post Banking / National Commercial Banks",
        "official_source_url": "https://www.indiapost.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Department of Posts and Ministry of Finance. High tax-free sovereign interest rate (currently 8.2% p.a.), exempt under EEE regime with Section 80C tax benefits."
    },
    {
        "service_id": "UDID_DISABILITY",
        "service_name": "Unique Disability ID (UDID) Card & Certificate",
        "service_category": "Disability Welfare & Inclusion",
        "intended_beneficiaries": "Persons with disabilities (PwDs) across all 21 categories recognized under the RPwD Act 2016",
        "basic_eligibility": "Indian citizen diagnosed with permanent or temporary disability of 40% or more confirmed by a certified Medical Board",
        "commonly_required_documents": "Aadhaar Card, Proof of Address, Passport photograph with visible impairment (if applicable), Existing medical diagnosis/hospital reports",
        "official_application_portal": "Swavlamban Portal (UDID)",
        "official_source_url": "https://www.swavlambancard.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Department of Empowerment of Persons with Disabilities (DEPwD). Single nationwide smart card eliminating the need to carry physical paper disability certificates for travel concessions and schemes."
    },
    {
        "service_id": "ATAL_PENSION_YOJANA",
        "service_name": "Atal Pension Yojana (APY)",
        "service_category": "Social Security & Pension",
        "intended_beneficiaries": "Workers in the unorganised sector seeking guaranteed lifetime monthly pension post-retirement",
        "basic_eligibility": "Indian citizen aged between 18 and 40 years; must possess savings bank account; not an income tax payer (as per revised 2022 guidelines)",
        "commonly_required_documents": "Savings Bank Account details, Aadhaar Card, Active mobile phone number for auto-debit SMS alerts",
        "official_application_portal": "myScheme / Protean CRA APY Portal",
        "official_source_url": "https://www.myscheme.gov.in/schemes/apy",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Pension Fund Regulatory and Development Authority (PFRDA). Guarantees monthly pension from ₹1,000 to ₹5,000 starting from age 60 until death, followed by spouse pension and return of corpus to nominee."
    },
    {
        "service_id": "IGNOAPS",
        "service_name": "Indira Gandhi National Old Age Pension Scheme (NSAP)",
        "service_category": "Senior Citizen Social Assistance",
        "intended_beneficiaries": "Indigent senior citizens aged 60 and above living below the poverty line (BPL)",
        "basic_eligibility": "Age 60 years or older; applicant's household must be classified as living Below Poverty Line (BPL) as per state/central criteria",
        "commonly_required_documents": "Age proof (Aadhaar / Voter ID / Birth record), BPL ration card / BPL certificate, Bank or Post Office savings passbook",
        "official_application_portal": "National Social Assistance Programme (NSAP) Portal",
        "official_source_url": "https://nsap.nic.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Ministry of Rural Development. Non-contributory central pension of ₹200/mo (age 60-79) and ₹500/mo (age 80+), supplemented by state government matching grants (ranging from ₹1,000 to ₹3,000/mo)."
    },
    {
        "service_id": "NCS_PORTAL",
        "service_name": "National Career Service (NCS) Registration",
        "service_category": "Employment & Career Advisory",
        "intended_beneficiaries": "Educated and uneducated job seekers, students, apprentices, and gig workers",
        "basic_eligibility": "Indian citizen aged 14 and above; seeking employment, vocational training, or career counselling",
        "commonly_required_documents": "Aadhaar / Voter ID / PAN, Educational qualification certificates / marksheets, Experience certificates (if applicable)",
        "official_application_portal": "National Career Service Portal",
        "official_source_url": "https://www.ncs.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Ministry of Labour & Employment. Provides free candidate profiles, nationwide Rozgar Mela invitations, career counselling, and job vacancy alerts. ₹0 registration fee."
    },
    {
        "service_id": "PM_SVANIDHI",
        "service_name": "PM Street Vendor's AtmaNirbhar Nidhi (PM SVANidhi)",
        "service_category": "Micro-Credit & Urban Livelihood",
        "intended_beneficiaries": "Urban and peri-urban street vendors and hawkers vending essential goods and services",
        "basic_eligibility": "Engaged in street vending in urban areas; possesses Certificate of Vending / ID Card issued by Urban Local Body (ULB) or Letter of Recommendation (LoR)",
        "commonly_required_documents": "Aadhaar Card, ULB Vending Card or LoR from Town Vending Committee (TVC), Bank account passbook",
        "official_application_portal": "PM SVANidhi Portal",
        "official_source_url": "https://pmsvanidhi.mohua.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Ministry of Housing and Urban Affairs (MoHUA). Collateral-free working capital loan: 1st tranche ₹10k, 2nd ₹20k, 3rd ₹50k with 7% interest subsidy and digital cashback up to ₹1,200/yr."
    },
    {
        "service_id": "PMKVY",
        "service_name": "Pradhan Mantri Kaushal Vikas Yojana (PMKVY 4.0)",
        "service_category": "Skill Training & Certification",
        "intended_beneficiaries": "School and college dropouts, uncertified artisans, and unemployed youth seeking marketable industry skills",
        "basic_eligibility": "Indian citizen aged 15 to 45; minimum educational criteria based on chosen National Skills Qualification Framework (NSQF) course",
        "commonly_required_documents": "Aadhaar Card, Bank account details, Educational proof / school leaving certificate, Passport photos",
        "official_application_portal": "Skill India Digital Hub",
        "official_source_url": "https://www.skillindiadigital.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Ministry of Skill Development and Entrepreneurship (MSDE). 100% government-funded short-term training, assessment, and certification with post-placement support."
    },
    {
        "service_id": "E_SHRAM",
        "service_name": "e-Shram National Database of Unorganised Workers",
        "service_category": "Unorganised Worker Social Protection",
        "intended_beneficiaries": "Unorganised sector workers including construction labourers, agricultural workers, gig/platform workers, street vendors, and domestic help",
        "basic_eligibility": "Age between 16 and 59 years; not a member of EPFO or ESIC; not an income tax payer",
        "commonly_required_documents": "Aadhaar Card, Aadhaar-linked active mobile number, Bank account number and IFSC",
        "official_application_portal": "e-Shram Portal",
        "official_source_url": "https://eshram.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Ministry of Labour & Employment. Issues 12-digit UAN card. Provides ₹2 lakh accidental death/permanent disability cover under PMSBY and acts as universal worker social security registry."
    },
    {
        "service_id": "BIRTH_CERTIFICATE_CRS",
        "service_name": "Birth Registration & Certificate (CRS)",
        "service_category": "Statutory Civil Registration",
        "intended_beneficiaries": "Newborn children and parents registering births occurring in Indian territory",
        "basic_eligibility": "Birth occurring in India; statutory reporting within 21 days for fee-exempt registration at local registrar (Gram Panchayat / Municipal Ward)",
        "commonly_required_documents": "Hospital/Institutional Birth Discharge Slip (Form 1), Parents' Aadhaar Cards, Address proof, Parents' Marriage Certificate (recommended)",
        "official_application_portal": "Civil Registration System Portal",
        "official_source_url": "https://crsorgi.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Registrar General of India (ORGI), Ministry of Home Affairs. Foundational legal identity document essential for school admissions, Aadhaar issuance, passport, and inheritance."
    },
    {
        "service_id": "POST_MATRIC_SCHOLARSHIP",
        "service_name": "Post-Matric Scholarship for SC/ST/OBC Students (NSP)",
        "service_category": "Higher Education Financial Aid",
        "intended_beneficiaries": "Students from SC, ST, and OBC communities enrolled in post-matriculation courses (Class 11, 12, ITI, Polytechnic, Degree, PG)",
        "basic_eligibility": "Belongs to notified SC/ST/OBC category; enrolled in recognised educational institution; annual family income <= ₹2.5 Lakh (SC/ST) or <= ₹1.5-₹2.5 Lakh (OBC as per state rules)",
        "commonly_required_documents": "Caste Certificate from competent revenue officer, Income Certificate (current financial year), Previous year marksheet, Bonafide student certificate / fee receipt, Aadhaar-seeded bank account",
        "official_application_portal": "National Scholarship Portal (NSP)",
        "official_source_url": "https://scholarships.gov.in/",
        "last_verified_date": "2026-10-25",
        "verification_notes": "Administered by Ministry of Social Justice & Empowerment and Ministry of Tribal Affairs via NSP. Covers mandatory non-refundable institutional fees plus monthly maintenance allowance via DBT."
    }
]

print("Building workbook...")

wb = openpyxl.Workbook()

# Setup sheets
ws_citizens = wb.active
ws_citizens.title = "Citizens"
ws_services = wb.create_sheet("Government Services")
ws_scenarios = wb.create_sheet("Eligibility Scenarios")
ws_guide = wb.create_sheet("Demo Guide")

# Enable gridlines on all sheets
for ws in [ws_citizens, ws_services, ws_scenarios, ws_guide]:
    ws.views.sheetView[0].showGridLines = True

# Professional styling definitions
font_title = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
font_subtitle = Font(name="Calibri", size=11, bold=True, color="475569")
font_notice = Font(name="Calibri", size=11, bold=True, color="991B1B")

fill_notice = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
fill_header_sub = PatternFill(start_color="0F2942", end_color="0F2942", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
font_data = Font(name="Calibri", size=10, color="0F172A")
font_bold_data = Font(name="Calibri", size=10, bold=True, color="0F172A")

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

thin_border_side = Side(border_style="thin", color="CBD5E1")
cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
header_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

# -------------------------------------------------------------
# POPULATE SHEET 1: Citizens (100 rows, 32 columns)
# -------------------------------------------------------------
CITIZEN_COLUMNS = [
    "citizen_id", "full_name", "age", "date_of_birth", "gender", "state", "district",
    "city_or_village", "rural_or_urban", "preferred_language", "education_level",
    "occupation", "employment_status", "annual_family_income_inr", "family_size",
    "marital_status", "has_children", "number_of_children", "disability_status",
    "student_status", "farmer_status", "senior_citizen_status", "demo_phone",
    "demo_email", "life_event", "life_event_date", "potential_government_services",
    "eligibility_notes", "required_documents", "application_channel", "verification_status",
    "demo_only"
]

ws_citizens.append(CITIZEN_COLUMNS)
ws_citizens.row_dimensions[1].height = 28

for col_idx in range(1, len(CITIZEN_COLUMNS) + 1):
    cell = ws_citizens.cell(row=1, column=col_idx)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = header_border

# Write 100 citizen data rows
csv_rows = []
for idx, r in enumerate(ALL_RECORDS):
    id_num = r[0]
    citizen_id = f"CIT-{id_num:04d}"
    full_name = r[1]
    dob_str = r[2]
    age = calc_age(dob_str)
    gender = r[3]
    state = r[4]
    district = r[5]
    city_village = r[6]
    rural_urban = r[7]
    lang = r[8]
    edu = r[9]
    occ = r[10]
    emp_status = r[11]
    income = int(r[12])
    fam_size = int(r[13])
    marital = r[14]
    has_children = bool(r[15])
    num_children = int(r[16])
    disability = r[17]
    student = bool(r[18])
    farmer = bool(r[19])
    senior = (age >= 60)
    phone = f"DEMO-PHONE-{id_num:04d}"
    slug = full_name.lower().replace(" ", ".").replace("'", "")
    email = f"{slug}.demo@example.com"
    life_event = r[20]
    life_event_date = r[21]
    pot_services = r[22]
    notes = r[23]
    req_docs = r[24]
    app_channel = r[25]
    ver_status = r[26]
    demo_only = True

    row_data = [
        citizen_id, full_name, age, dob_str, gender, state, district,
        city_village, rural_urban, lang, edu, occ, emp_status,
        income, fam_size, marital, has_children, num_children,
        disability, student, farmer, senior, phone, email,
        life_event, life_event_date, pot_services, notes, req_docs,
        app_channel, ver_status, demo_only
    ]
    csv_rows.append(row_data)
    ws_citizens.append(row_data)

    row_num = idx + 2
    ws_citizens.row_dimensions[row_num].height = 20
    is_even = (row_num % 2 == 0)
    row_fill = fill_zebra if is_even else fill_white

    for col_idx in range(1, len(CITIZEN_COLUMNS) + 1):
        cell = ws_citizens.cell(row=row_num, column=col_idx)
        cell.font = font_data
        cell.fill = row_fill
        cell.border = cell_border

        # Alignments & formats
        if col_idx in [1, 3, 4, 5, 9, 15, 17, 18, 20, 21, 22, 23, 26, 31, 32]:
            cell.alignment = align_center
        elif col_idx == 14:
            cell.alignment = align_right
            cell.number_format = "₹#,##,##0"
        else:
            cell.alignment = align_left

ws_citizens.freeze_panes = "A2"
ws_citizens.auto_filter.ref = f"A1:{get_column_letter(len(CITIZEN_COLUMNS))}101"

# -------------------------------------------------------------
# POPULATE SHEET 2: Government Services (18 rows, 10 columns)
# -------------------------------------------------------------
SERVICE_COLUMNS = [
    "service_id", "service_name", "service_category", "intended_beneficiaries",
    "basic_eligibility", "commonly_required_documents", "official_application_portal",
    "official_source_url", "last_verified_date", "verification_notes"
]

ws_services.append(SERVICE_COLUMNS)
ws_services.row_dimensions[1].height = 28

for col_idx in range(1, len(SERVICE_COLUMNS) + 1):
    cell = ws_services.cell(row=1, column=col_idx)
    cell.font = font_header
    cell.fill = fill_header_sub
    cell.alignment = align_center
    cell.border = header_border

for idx, s in enumerate(SERVICES_DATA):
    row_num = idx + 2
    row_data = [
        s["service_id"], s["service_name"], s["service_category"],
        s["intended_beneficiaries"], s["basic_eligibility"],
        s["commonly_required_documents"], s["official_application_portal"],
        s["official_source_url"], s["last_verified_date"], s["verification_notes"]
    ]
    ws_services.append(row_data)
    ws_services.row_dimensions[row_num].height = 36
    is_even = (row_num % 2 == 0)
    row_fill = fill_zebra if is_even else fill_white

    for col_idx in range(1, len(SERVICE_COLUMNS) + 1):
        cell = ws_services.cell(row=row_num, column=col_idx)
        cell.font = font_data
        cell.fill = row_fill
        cell.border = cell_border
        if col_idx in [1, 9]:
            cell.alignment = align_center
        else:
            cell.alignment = align_left

ws_services.freeze_panes = "A2"
ws_services.auto_filter.ref = f"A1:{get_column_letter(len(SERVICE_COLUMNS))}{len(SERVICES_DATA)+1}"

# -------------------------------------------------------------
# POPULATE SHEET 3: Eligibility Scenarios (100 rows, 9 columns)
# -------------------------------------------------------------
SCENARIO_COLUMNS = [
    "citizen_id", "scenario_id", "life_event", "potential_service",
    "eligibility_status", "reason", "missing_information",
    "next_recommended_step", "official_source_url"
]

ws_scenarios.append(SCENARIO_COLUMNS)
ws_scenarios.row_dimensions[1].height = 28

for col_idx in range(1, len(SCENARIO_COLUMNS) + 1):
    cell = ws_scenarios.cell(row=1, column=col_idx)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = header_border

# Status specific badge fills for visual clarity
fill_potential = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")     # Light green
fill_not_eligible = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # Light red
fill_needs_verif = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")   # Light amber

font_green = Font(name="Calibri", size=10, bold=True, color="166534")
font_red = Font(name="Calibri", size=10, bold=True, color="991B1B")
font_amber = Font(name="Calibri", size=10, bold=True, color="92400E")

for idx, r in enumerate(ALL_RECORDS):
    id_num = r[0]
    citizen_id = f"CIT-{id_num:04d}"
    scenario_id = f"SCN-{id_num:04d}"
    life_event = r[20]
    potential_service = r[27]
    elig_status = r[28]
    reason = r[29]
    missing_info = r[30]
    next_step = r[31]
    source_url = r[32]

    # Harmonize APY url to official myScheme .gov.in URL
    if "npscra.nsdl.co.in" in source_url:
        source_url = "https://www.myscheme.gov.in/schemes/apy"

    row_data = [
        citizen_id, scenario_id, life_event, potential_service,
        elig_status, reason, missing_info, next_step, source_url
    ]
    ws_scenarios.append(row_data)

    row_num = idx + 2
    ws_scenarios.row_dimensions[row_num].height = 24
    is_even = (row_num % 2 == 0)
    row_fill = fill_zebra if is_even else fill_white

    for col_idx in range(1, len(SCENARIO_COLUMNS) + 1):
        cell = ws_scenarios.cell(row=row_num, column=col_idx)
        cell.font = font_data
        cell.fill = row_fill
        cell.border = cell_border

        if col_idx in [1, 2]:
            cell.alignment = align_center
        elif col_idx == 5:
            cell.alignment = align_center
            if elig_status == "POTENTIAL":
                cell.fill = fill_potential
                cell.font = font_green
            elif elig_status == "NOT_ELIGIBLE":
                cell.fill = fill_not_eligible
                cell.font = font_red
            else:
                cell.fill = fill_needs_verif
                cell.font = font_amber
        else:
            cell.alignment = align_left

ws_scenarios.freeze_panes = "A2"
ws_scenarios.auto_filter.ref = f"A1:{get_column_letter(len(SCENARIO_COLUMNS))}101"

# -------------------------------------------------------------
# POPULATE SHEET 4: Demo Guide
# -------------------------------------------------------------
GUIDE_CONTENT = [
    ("SYNTHETIC DEMO DATA NOTICE — STRICT PROTOTYPE SAFETY", [
        ("WARNING", "THIS WORKBOOK CONTAINS 100% FICTIONAL SYNTHETIC PROFILES CREATED STRICTLY FOR TIMBRE SYSTEM TESTING, SIMULATION, AND DEMONSTRATIONS."),
        ("CRITICAL INVARIANT", "NEVER USE THIS DATA TO CONTACT REAL PEOPLE OR SUBMIT REAL GOVERNMENT APPLICATIONS."),
        ("SAFETY CONTROLS", "1. All phone numbers use reserved non-dialable placeholders: DEMO-PHONE-0001 through DEMO-PHONE-0100.\n"
                            "2. All email addresses are fictional domains under @example.com.\n"
                            "3. No real Aadhaar, PAN, Voter ID, or bank account numbers are generated or present.\n"
                            "4. Every record contains DEMO_ONLY = TRUE as an immutable boolean gatekeeper.\n"
                            "5. Telephony adapters in Timbre enforce a strict mock mode when DEMO_ONLY = TRUE, preventing any outbound carrier traffic.")
    ]),
    ("WORKBOOK STRUCTURE & SCHEMA OVERVIEW", [
        ("Sheet 1: Citizens", "Contains exactly 100 fictional citizens across 24 Indian states and UTs with 32 demographic and civic columns.\n"
                              "Demographics span teenagers (16-17), first-time 18-year-old voters, college students, job seekers, gig workers, smallholder farmers, low-income BPL households, newlyweds, newborn parents, synthetic disability scenarios, pre-retirees (55-59), senior citizens (60-69), and super-seniors (70-85)."),
        ("Sheet 2: Government Services", "A curated catalogue of 18 verified Indian statutory schemes and civic services across 10 structured columns.\n"
                                        "All official portal URLs are strictly verified .gov.in or .nic.in sources with zero invented URLs or benefits."),
        ("Sheet 3: Eligibility Scenarios", "100 distinct evaluation scenarios directly mapped 1-to-1 with CIT-0001 through CIT-0100.\n"
                                          "Classifies eligibility as POTENTIAL (67%), NEEDS_VERIFICATION (18%), or NOT_ELIGIBLE (15%) with statutory rationales and concrete civic next steps."),
        ("Sheet 4: Demo Guide", "This documentation sheet explaining architecture, simulation mechanics, backend loading, and migration path.")
    ]),
    ("HOW TO CONDUCT TIMBRE DEMONSTRATIONS", [
        ("Demo Flow 1: Turning 18 Milestone", "1. Select citizen CIT-0001 (Aarav Sharma) or CIT-0002 (Priya Patel).\n"
                                              "2. In the Flutter admin simulation cockpit, change the date of birth to cross the 18th birthday threshold.\n"
                                              "3. Observe the Eligibility Engine trigger an automated evaluation for VOTER_ID_ECI (Form 6) and LEARNERS_LICENCE_MORTH.\n"
                                              "4. Inspect the Call Monitor screen to observe the deduplicated simulated outbound call job generated with idempotency key CALL_AGE18_CIT-0001_2026.\n"
                                              "5. View the simulated Hinglish voice transcript explaining Form 6 registration and Parivahan e-KYC."),
        ("Demo Flow 2: Smallholder Farmer Advisory", "1. Select citizen CIT-0046 (Vikram Singh Rathore) or CIT-0050 (Naresh Meena).\n"
                                                    "2. Simulate the sowing season life event on the dashboard.\n"
                                                    "3. The engine checks landholding records (Jamabandi/Khatauni) and identifies PM-KISAN ₹6,000 DBT and subsidized Kisan Credit Card (KCC) at 4% interest.\n"
                                                    "4. Review the generated checklist of documents needed for Tehsil/CSC verification."),
        ("Demo Flow 3: Super-Senior Health Universal Cover", "1. Select citizen CIT-0095 (Rajeshwar Rao, Age 72) or CIT-0096 (Kamala Bai, Age 78).\n"
                                                            "2. Trigger the 70th birthday milestone event.\n"
                                                            "3. The engine verifies the attainment of age 70 and identifies the Ayushman Vaya Vandana Card (PM-JAY 70+), which provides universal ₹5 Lakh free hospital cover regardless of family income.\n"
                                                            "4. Demonstrate how the Hinglish voice assistant guides the senior to register on beneficiary.nha.gov.in using Aadhaar OTP.")
    ]),
    ("BACKEND INTEGRATION GUIDE (FASTAPI / POSTGRESQL)", [
        ("Step 1: CSV Loading via Python / Pandas", "import pandas as pd\n"
                                                   "df = pd.read_csv('data/demo/Timbre_Fake_Citizens.csv')\n"
                                                   "# Safety guard check\n"
                                                   "assert (df['demo_only'] == True).all(), 'Safety violation: Non-demo records found!'\n"
                                                   "print(f'Successfully loaded {len(df)} synthetic citizens.')"),
        ("Step 2: PostgreSQL Seed Ingestion", "COPY citizens(citizen_id, full_name, age, date_of_birth, gender, state, district, ...)\n"
                                             "FROM '/absolute/path/to/Timbre_Fake_Citizens.csv' WITH (FORMAT csv, HEADER true);\n"
                                             "-- Always verify the is_fictional / demo_only safety constraint:\n"
                                             "ALTER TABLE citizens ADD CONSTRAINT chk_demo_safety CHECK (demo_only = TRUE);"),
        ("Step 3: Eligibility Rule Engine Pipeline", "1. Demographic Filter: Match age (min_age <= citizen.age <= max_age) and gender.\n"
                                                    "2. Income Ceiling Check: For welfare schemes (PM-JAY, PMAY-G, NSP), verify annual_family_income_inr <= threshold.\n"
                                                    "3. Occupational & Land Filters: Check farmer_status for PM-KISAN/KCC; student_status for scholarships.\n"
                                                    "4. Deduplication & Idempotency: Generate job key: CALL_{SERVICE}_{CITIZEN_ID}_{YEAR}. If key exists in DB, discard duplicate event.\n"
                                                    "5. Flag Cases Requiring Verification: If missing essential records (e.g. land mutation or medical board evaluation), output NEEDS_VERIFICATION instead of auto-enrolling.")
    ]),
    ("SAFETY ARCHITECTURE AGAINST REAL-WORLD DIALING", [
        ("Invariable Guard 1: Non-Dialable MSISDNs", "All phone numbers strictly use the pattern DEMO-PHONE-xxxx. In standard telecommunication networks (ITU-T E.164), non-numeric alphabetic strings cannot be resolved or routed by SIP trunk providers or cellular base stations."),
        ("Invariable Guard 2: Backend Mock Adapter", "The Timbre Voice Orchestrator checks citizen.demo_only == True before invoking any CPaaS adapter (Exotel/Twilio/Bolna). If True, the request is routed exclusively to the local In-Memory Mock Voice Pipeline."),
        ("Invariable Guard 3: Immutable Database Constraints", "A PostgreSQL CHECK constraint (demo_only = TRUE) guarantees that real citizen records can never be accidentally inserted into development tables.")
    ]),
    ("MIGRATION PATH TO CONSENTED PILOT TESTING", [
        ("Phase 1: Synthetic Sandbox (Current)", "Full end-to-end testing of conversational dialogue, turn-taking, and rule logic using this 100-citizen synthetic dataset."),
        ("Phase 2: Closed Internal Consented Cohort", "Transition to a small pilot cohort (e.g., 20 volunteers) under signed Informed Consent Forms adhering to India's Digital Personal Data Protection (DPDP) Act 2023.\n"
                                                     "- Provide clear notice regarding voice data processing and purpose limitation.\n"
                                                     "- Store Aadhaar numbers exclusively inside UIDAI-compliant Aadhaar Data Vaults with tokenized reference IDs.\n"
                                                     "- Enforce an immediate right to data erasure (DPDP Section 12)."),
        ("Phase 3: TRAI / DoT Compliant Telephony", "Prior to any actual PSTN/cellular voice dialing, acquire authorized 140-series telemarketing/advisory CLI allocation, register DLT templates under TRAI regulations, and integrate explicit opt-in verification.")
    ])
]

# Write Demo Guide
guide_row = 1
ws_guide.row_dimensions[guide_row].height = 32
cell = ws_guide.cell(row=guide_row, column=1, value="TIMBRE PROJECT — SYNTHETIC DEMO DATASET COMPREHENSIVE GUIDE")
cell.font = font_title
cell.alignment = align_left
guide_row += 1

cell = ws_guide.cell(row=guide_row, column=1, value="NOTICE: SYNTHETIC DEMO DATA — NOT FOR REAL-WORLD APPLICATION OR OFFICIAL GOVERNMENT DECISIONS")
cell.font = font_notice
cell.fill = fill_notice
cell.alignment = align_left
ws_guide.row_dimensions[guide_row].height = 24
guide_row += 2

for section_title, entries in GUIDE_CONTENT:
    ws_guide.cell(row=guide_row, column=1, value=section_title).font = font_subtitle
    ws_guide.row_dimensions[guide_row].height = 24
    guide_row += 1

    for label, text in entries:
        c1 = ws_guide.cell(row=guide_row, column=1, value=label)
        c1.font = font_bold_data
        c1.alignment = Alignment(horizontal="left", vertical="top")
        c1.border = cell_border

        c2 = ws_guide.cell(row=guide_row, column=2, value=text)
        c2.font = font_data
        c2.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c2.border = cell_border

        # Estimate row height based on text linebreaks
        line_count = text.count("\n") + 1
        ws_guide.row_dimensions[guide_row].height = max(24, line_count * 18 + 8)
        guide_row += 1

    guide_row += 1

ws_guide.column_dimensions["A"].width = 30
ws_guide.column_dimensions["B"].width = 110

# Auto-adjust column widths on data sheets
def auto_fit_columns(ws, min_width=12, max_width=50, padding=4):
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val = cell.value
            if val is not None:
                lines = str(val).split("\n")
                for line in lines:
                    if len(line) > max_len:
                        max_len = len(line)
        ws.column_dimensions[col_letter].width = min(max(max_len + padding, min_width), max_width)

print("Formatting column widths...")
auto_fit_columns(ws_citizens, min_width=12, max_width=45, padding=3)
auto_fit_columns(ws_services, min_width=15, max_width=55, padding=3)
auto_fit_columns(ws_scenarios, min_width=14, max_width=55, padding=3)

# Add dropdown data validation for status columns
print("Adding Excel data validation rules...")
dv_gender = DataValidation(type="list", formula1='"MALE,FEMALE,TRANSGENDER,NON-BINARY"', allow_blank=True)
dv_rural_urban = DataValidation(type="list", formula1='"RURAL,URBAN"', allow_blank=True)
dv_bool = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
dv_verif = DataValidation(type="list", formula1='"VERIFIED_MOCK,PENDING_DOCS,ELIGIBLE_PROVISIONAL"', allow_blank=True)
dv_demo = DataValidation(type="list", formula1='"TRUE"', allow_blank=True)

ws_citizens.add_data_validation(dv_gender)
ws_citizens.add_data_validation(dv_rural_urban)
ws_citizens.add_data_validation(dv_bool)
ws_citizens.add_data_validation(dv_verif)
ws_citizens.add_data_validation(dv_demo)

dv_gender.add("E2:E101")
dv_rural_urban.add("I2:I101")
dv_bool.add("Q2:Q101")
dv_bool.add("T2:T101")
dv_bool.add("U2:U101")
dv_bool.add("V2:V101")
dv_verif.add("AE2:AE101")
dv_demo.add("AF2:AF101")

dv_scenario_status = DataValidation(type="list", formula1='"POTENTIAL,NOT_ELIGIBLE,NEEDS_VERIFICATION"', allow_blank=True)
ws_scenarios.add_data_validation(dv_scenario_status)
dv_scenario_status.add("E2:E101")

# Ensure target directories exist
target_dirs = [
    "/Users/sidh._./Timbre/data/demo",
    "/Users/sidh._./Timbre/FakeData/data/demo",
    "/Users/sidh._./Timbre/FakeData"
]
for d in target_dirs:
    os.makedirs(d, exist_ok=True)

excel_repo_path = "/Users/sidh._./Timbre/data/demo/Timbre_Fake_Citizens.xlsx"
csv_repo_path = "/Users/sidh._./Timbre/data/demo/Timbre_Fake_Citizens.csv"

print(f"Saving Excel workbook to {excel_repo_path}...")
wb.save(excel_repo_path)

print(f"Saving CSV file to {csv_repo_path}...")
with open(csv_repo_path, "w", newline="", encoding="utf-8") as f:
    writer = csv.writer(f)
    writer.writerow(CITIZEN_COLUMNS)
    writer.writerows(csv_rows)

# Duplicate to FakeData workspace for convenience
for d in ["/Users/sidh._./Timbre/FakeData/data/demo", "/Users/sidh._./Timbre/FakeData"]:
    shutil.copy2(excel_repo_path, os.path.join(d, "Timbre_Fake_Citizens.xlsx"))
    shutil.copy2(csv_repo_path, os.path.join(d, "Timbre_Fake_Citizens.csv"))

print("Dataset generation completed successfully!")
