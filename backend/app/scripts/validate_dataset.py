"""Validate the fictional demo CSV and write an aggregate-only report."""
import csv
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Dict, Optional
from urllib.parse import urlparse

import openpyxl

from app.config import settings
from app.services.catalogue_data import OFFICIAL_SERVICES_CATALOGUE

REPO_ROOT = Path(__file__).resolve().parents[3]
CSV_PATH = REPO_ROOT / "data" / "demo" / "Timbre_Fake_Citizens.csv"
XLSX_PATH = REPO_ROOT / "data" / "demo" / "Timbre_Fake_Citizens.xlsx"
REPORT_PATH = REPO_ROOT / "data" / "reports" / "dataset_validation.json"
REQUIRED_FIELDS = {
    "citizen_id", "first_name", "last_name", "name", "date_of_birth", "age",
    "gender", "state", "district", "synthetic_family_income", "family_size",
    "children", "is_student", "is_farmer", "is_senior_citizen", "demo_only",
    "phone_number", "email", "potentially_relevant_services", "required_documents",
}
INTEGER_FIELDS = {"age", "synthetic_family_income", "family_size", "children"}
BOOLEAN_FIELDS = {"is_student", "is_farmer", "is_senior_citizen", "demo_only"}
ARRAY_FIELDS = {"potentially_relevant_services", "required_documents"}
BOOL_VALUES = {"true", "false", "1", "0"}
OFFICIAL_HOST_SUFFIXES = (".gov.in", ".nic.in", ".eci.gov.in", ".gov")


def validate_dataset(
    csv_path: Path = CSV_PATH,
    xlsx_path: Path = XLSX_PATH,
    report_path: Optional[Path] = REPORT_PATH,
    reference_date: Optional[date] = None,
) -> Dict[str, object]:
    reference_date = reference_date or settings.parsed_reference_date
    errors: Counter = Counter()
    warnings: Counter = Counter()
    rows = []
    headers = []
    raw = csv_path.read_bytes()
    with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        headers = reader.fieldnames or []
        rows = list(reader)

    missing_columns = REQUIRED_FIELDS - set(headers)
    if missing_columns:
        errors["required_columns_missing"] += len(missing_columns)

    identifiers = [row.get("citizen_id", "").strip() for row in rows]
    duplicate_ids = len(identifiers) - len(set(identifiers))
    if duplicate_ids:
        errors["duplicate_identifiers"] += duplicate_ids

    valid_service_ids = {service["id"] for service in OFFICIAL_SERVICES_CATALOGUE}
    mismatched_age = 0
    personal_data_flags = 0
    bad_types = 0
    missing_values = 0
    malformed_ids = 0
    non_synthetic_markers = 0
    invalid_lists = 0
    invalid_services = 0
    invalid_dates = 0
    seen_contact = set()

    for row in rows:
        for field in REQUIRED_FIELDS:
            if not (row.get(field) or "").strip():
                missing_values += 1
        cid = row.get("citizen_id", "").strip()
        if not re.fullmatch(r"CIT-[A-Z0-9-]+", cid):
            malformed_ids += 1
        if row.get("demo_only", "").lower() not in {"true", "1"}:
            non_synthetic_markers += 1
        if not row.get("phone_number", "").startswith("DEMO-PHONE-") or not row.get("email", "").lower().endswith("@example.com"):
            non_synthetic_markers += 1
        for field in INTEGER_FIELDS:
            try:
                int(row.get(field, ""))
            except (TypeError, ValueError):
                bad_types += 1
        for field in BOOLEAN_FIELDS:
            if row.get(field, "").lower() not in BOOL_VALUES:
                bad_types += 1
        try:
            dob = date.fromisoformat(row.get("date_of_birth", ""))
            if dob > reference_date:
                invalid_dates += 1
            computed_age = reference_date.year - dob.year - ((reference_date.month, reference_date.day) < (dob.month, dob.day))
            if computed_age != int(row.get("age", "")):
                mismatched_age += 1
        except (TypeError, ValueError):
            invalid_dates += 1
        for field in ARRAY_FIELDS:
            try:
                value = json.loads(row.get(field, ""))
                if not isinstance(value, list):
                    invalid_lists += 1
                elif field == "potentially_relevant_services" and any(item not in valid_service_ids for item in value):
                    invalid_services += 1
            except (TypeError, ValueError):
                invalid_lists += 1
        for value in row.values():
            text = str(value or "")
            digits = re.sub(r"\D", "", text)
            if len(digits) == 12 or re.search(r"(?i)\b[A-Z]{5}[0-9]{4}[A-Z]\b", text):
                personal_data_flags += 1
            if re.search(r"(?i)(nvapi-|sk-[a-z0-9]{12,}|api[_-]?key\s*=)", text):
                personal_data_flags += 1
        contact = (row.get("phone_number", ""), row.get("email", ""))
        if contact in seen_contact:
            warnings["duplicate_placeholder_contacts"] += 1
        seen_contact.add(contact)

    for label, count in {
        "required_cells_missing": missing_values,
        "invalid_identifier_format": malformed_ids,
        "non_synthetic_contact_or_marker": non_synthetic_markers,
        "invalid_field_types": bad_types,
        "invalid_or_future_dates": invalid_dates,
        "age_dob_mismatch": mismatched_age,
        "invalid_json_arrays": invalid_lists,
        "unknown_service_codes": invalid_services,
        "possible_personal_data_or_secret": personal_data_flags,
    }.items():
        if count:
            errors[label] += count

    workbook_rows = None
    workbook_sheets = []
    workbook_field_count = None
    try:
        workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        workbook_sheets = workbook.sheetnames
        data_sheet = workbook["Fake_Citizens_100"]
        workbook_rows = max(0, data_sheet.max_row - 1)
        workbook_headers = [str(value or "").strip() for value in next(data_sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        workbook_field_count = len(workbook_headers)
        canonical_header = {"id": "citizen_id", "full_name": "name", "synthetic_family_income_inr": "synthetic_family_income"}
        normalized_xlsx_headers = {canonical_header.get(header, header) for header in workbook_headers}
        normalized_csv_headers = {canonical_header.get(header, header) for header in headers}
        if normalized_xlsx_headers != normalized_csv_headers:
            errors["csv_xlsx_header_mismatch"] += 1
        if workbook_rows != len(rows):
            errors["csv_xlsx_record_count_mismatch"] += abs(workbook_rows - len(rows)) or 1
        workbook.close()
    except (OSError, KeyError, ValueError):
        errors["xlsx_missing_or_invalid"] += 1

    invalid_catalogue_sources = 0
    for service in OFFICIAL_SERVICES_CATALOGUE:
        parsed = urlparse(service.get("official_source_url", ""))
        host = (parsed.hostname or "").lower()
        if parsed.scheme != "https" or not any(host.endswith(suffix) or host == suffix.lstrip(".") for suffix in OFFICIAL_HOST_SUFFIXES):
            invalid_catalogue_sources += 1
    if invalid_catalogue_sources:
        errors["non_official_or_invalid_source_url"] += invalid_catalogue_sources
    warnings["source_content_needs_human_review"] += 1
    pending_catalogue_reviews = sum(service.get("verification_status") != "VERIFIED" for service in OFFICIAL_SERVICES_CATALOGUE)
    if pending_catalogue_reviews:
        warnings["catalogue_entries_need_human_review"] += pending_catalogue_reviews

    report: Dict[str, object] = {
        "generated_at_utc": datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
        "dataset": "Timbre_Fake_Citizens.csv",
        "synthetic_dataset_assertion": "All records carry demo_only and placeholder contact markers; this validation cannot independently prove how source data was created.",
        "reference_date": reference_date.isoformat(),
        "csv_sha256": hashlib.sha256(raw).hexdigest(),
        "csv_record_count": len(rows),
        "csv_field_count": len(headers),
        "xlsx_record_count": workbook_rows,
        "xlsx_field_count": workbook_field_count,
        "xlsx_column_aliases_checked": ["id/citizen_id", "full_name/name", "synthetic_family_income_inr/synthetic_family_income"],
        "xlsx_sheets": workbook_sheets,
        "duplicate_identifier_count": duplicate_ids,
        "errors": dict(sorted(errors.items())),
        "error_count": sum(errors.values()),
        "warnings": dict(sorted(warnings.items())),
        "warning_count": sum(warnings.values()),
        "catalogue_review_pending_count": pending_catalogue_reviews,
        "action_taken": "Original CSV and XLSX preserved. No source rows were changed. Report contains aggregate counts only.",
        "eligibility_source_limit": "Government source URLs are checked for HTTPS and government-domain form only. Eligibility rules, fees, required documents, deadlines, and links must be manually checked against the current official sources before relying on them.",
    }
    report["status"] = "PASS" if not errors else "FAIL"
    if report_path:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    return report


if __name__ == "__main__":
    result = validate_dataset()
    print(json.dumps({key: result[key] for key in ("status", "csv_record_count", "csv_field_count", "xlsx_record_count", "duplicate_identifier_count", "errors", "error_count", "warnings", "warning_count")}, indent=2))
